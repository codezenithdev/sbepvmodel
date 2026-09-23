"""Editable Word rendering of exactly the same verified report blocks as PDF."""
import base64
from datetime import datetime, timezone
import hashlib
from io import BytesIO
import re

from sbepv.technoeconomic_math import equation_runs


def _add_equation_runs(paragraph, source, *, size=None, mono=True):
    """Add runs to a paragraph, raising/lowering sub/superscripts.

    ``mono`` renders displayed equations in a monospaced face; inline variable
    mentions in prose keep the surrounding body font. Newlines inside the
    equation become explicit line breaks so multi-line equations keep their
    layout inside a table cell.
    """
    from docx.shared import Pt
    for text, script in equation_runs(source):
        segments = text.split('\n')
        for offset, segment in enumerate(segments):
            if offset:
                paragraph.add_run().add_break()
            if not segment:
                continue
            run = paragraph.add_run(segment)
            if mono:
                run.font.name = 'Consolas'
            if size is not None:
                run.font.size = Pt(size)
            if script == 'super':
                run.font.superscript = True
            elif script == 'sub':
                run.font.subscript = True
    return paragraph


def _figure_bookmark(figure_id):
    """Stable Word bookmark names (letter first, no spaces, at most 40 chars)."""
    value=str(figure_id)
    return 'fig_'+re.sub(r'[^A-Za-z0-9_]', '_', value)[:26]+'_'+hashlib.sha256(value.encode()).hexdigest()[:8]


def _table_bookmark(table_id):
    return 'tbl_' + _figure_bookmark(table_id)[4:]


def _footer_timestamp(report):
    value=report.get('analysis_at') or report.get('generated_at')
    label='Completed' if report.get('analysis_at') else 'Exported'
    if not value:
        return 'Analysis timestamp unavailable'
    try:
        parsed=datetime.fromisoformat(str(value).replace('Z','+00:00'))
        if parsed.tzinfo is not None:
            parsed=parsed.astimezone(timezone.utc)
            return f'{label} {parsed:%Y-%m-%d %H:%M} UTC'
        return f'{label} {parsed:%Y-%m-%d %H:%M}'
    except ValueError:
        return f'{label} {value}'


def render_docx(report):
    """Build raw editable OOXML; exported files must pass refresh_docx_fields."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.shared import Inches, Pt, RGBColor

    document=Document()
    for border in document.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    section=document.sections[0]
    section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.left_margin=section.right_margin=Pt(45)
    section.top_margin=Pt(42);section.bottom_margin=Pt(58)
    section.header_distance=Pt(19);section.footer_distance=Pt(14)
    for style_name,size in (('Normal',11),('Title',22),('Subtitle',11),('Heading 1',17),('Heading 2',11.5),('Heading 3',11),('Caption',8)):
        style=document.styles[style_name]
        style.font.name='Arial';style.font.size=Pt(size)
        style.font.color.rgb=RGBColor.from_string('404040' if style_name in ('Heading 2','Heading 3') else '333333')
        style.font.bold=style_name in ('Heading 2','Heading 3')
        style.paragraph_format.space_after=Pt(7)
        style.paragraph_format.line_spacing=1.1
        if style_name.startswith('Heading') or style_name=='Title':
            style.paragraph_format.keep_with_next=True
    document.styles['Normal'].paragraph_format.widow_control=True
    document.styles['Heading 1'].paragraph_format.space_before=Pt(8)
    document.styles['Title'].paragraph_format.space_before=Pt(0)
    outline=OxmlElement('w:outlineLvl');outline.set(qn('w:val'),'9')
    document.styles['Title'].element.get_or_add_pPr().append(outline)
    document.core_properties.title=report['title']
    document.core_properties.author='SBE PV Dashboard'
    document.core_properties.subject='Calibration annual simulation and technoeconomic comparison'
    document.core_properties.version=report['version']
    # Link the heading styles to a real multilevel list so edits in Word
    # renumber headings and the native contents field automatically.
    numbering=document.part.numbering_part.element
    abstract_id=max((int(node.get(qn('w:abstractNumId'))) for node in numbering.findall(qn('w:abstractNum'))),default=-1)+1
    num_id=max((int(node.get(qn('w:numId'))) for node in numbering.findall(qn('w:num'))),default=0)+1
    abstract=OxmlElement('w:abstractNum');abstract.set(qn('w:abstractNumId'),str(abstract_id))
    multilevel=OxmlElement('w:multiLevelType');multilevel.set(qn('w:val'),'multilevel');abstract.append(multilevel)
    for level in range(3):
        definition=OxmlElement('w:lvl');definition.set(qn('w:ilvl'),str(level))
        for tag,value in (('start','1'),('numFmt','decimal'),('pStyle',f'Heading{level+1}'),('suff','space'),('lvlText','.'.join('%'+str(i+1) for i in range(level+1))),('lvlJc','left')):
            item=OxmlElement('w:'+tag);item.set(qn('w:val'),value);definition.append(item)
        props=OxmlElement('w:pPr');indent=OxmlElement('w:ind')
        indent.set(qn('w:left'),'0');indent.set(qn('w:hanging'),'0');props.append(indent);definition.append(props)
        abstract.append(definition)
    # Abstract definitions precede concrete instances in numbering.xml.
    first_num=numbering.find(qn('w:num'))
    numbering.insert(numbering.index(first_num) if first_num is not None else len(numbering),abstract)
    concrete=OxmlElement('w:num');concrete.set(qn('w:numId'),str(num_id))
    reference=OxmlElement('w:abstractNumId');reference.set(qn('w:val'),str(abstract_id));concrete.append(reference);numbering.append(concrete)
    for level in range(3):
        props=document.styles[f'Heading {level+1}'].element.get_or_add_pPr().get_or_add_numPr()
        props.get_or_add_ilvl().val=level;props.get_or_add_numId().val=num_id
    def field(paragraph,instruction,placeholder=''):
        run=paragraph.add_run()
        begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin');begin.set(qn('w:dirty'),'true')
        code=OxmlElement('w:instrText');code.set(qn('xml:space'),'preserve');code.text=' '+instruction+' '
        separate=OxmlElement('w:fldChar');separate.set(qn('w:fldCharType'),'separate')
        result=OxmlElement('w:t');result.text=str(placeholder)
        end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end')
        for item in (begin,code,separate,result,end):run._r.append(item)
        return run
    # Word's built-in Footer style has a center tab stop. Remove it instead
    # of letting the page label stop in the middle of the page.
    footer_style=document.styles['Footer']
    footer_style.paragraph_format.tab_stops.clear_all()
    footer_style.paragraph_format.tab_stops.add_tab_stop(Pt(522),WD_TAB_ALIGNMENT.RIGHT)
    footer_style.paragraph_format.space_before=Pt(0);footer_style.paragraph_format.space_after=Pt(0)
    footer_style.paragraph_format.line_spacing=Pt(10)
    footer_style.font.name='Arial';footer_style.font.size=Pt(8.5)
    name=section.footer.paragraphs[0]
    name.add_run(str(report.get('analysis_name') or 'PV comparison'))
    footer=section.footer.add_paragraph(style='Footer')
    footer.add_run(_footer_timestamp(report)+'\t')
    footer.add_run('Page ');field(footer,'PAGE','1');footer.add_run(' of ');field(footer,'NUMPAGES','1')
    for paragraph in section.footer.paragraphs:
        paragraph.alignment=WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.keep_with_next=False
        for run in paragraph.runs:run.font.size=Pt(8.5)
    update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');document.settings.element.append(update)
    bookmark_counter=0
    def bookmark_start(paragraph,name):
        nonlocal bookmark_counter
        identity=str(bookmark_counter);bookmark_counter+=1
        start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),identity);start.set(qn('w:name'),name)
        paragraph._p.append(start)
        return identity
    def bookmark_end(paragraph,identity):
        end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),identity);paragraph._p.append(end)
    for block in report['blocks']:
        kind=block['kind']
        if kind=='pagebreak':
            p=document.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0)
            p.add_run().add_break(WD_BREAK.PAGE)
        elif kind=='title':
            document.add_paragraph(block['text'],style='Title')
            if block.get('subtitle'):
                # Word's Subtitle style carries the 11pt detail size; the cover
                # subheader needs its own step between the title and the run details.
                subtitle=document.add_paragraph(style='Subtitle')
                subtitle.add_run(block['subtitle']).font.size=Pt(14)
                subtitle.paragraph_format.keep_with_next=True
        elif kind=='heading':
            p=document.add_paragraph(style=f"Heading {block['level']}")
            p.paragraph_format.page_break_before=block.get('page') is True
            props=p._p.get_or_add_pPr().get_or_add_numPr()
            props.get_or_add_ilvl().val=block['level']-1
            props.get_or_add_numId().val=num_id if block.get('number') else 0
            identity=bookmark_start(p,block['anchor'].replace('-','_'))
            p.add_run(block['text']);bookmark_end(p,identity)
        elif kind=='paragraph':
            style=block.get('style','body')
            p=document.add_paragraph(style='Subtitle' if style=='subtitle' else 'Normal')
            if block.get('math'):
                _add_equation_runs(p,block['text'])
            elif block.get('inline_math'):
                _add_equation_runs(p,block['text'],mono=False)
            else:
                for segment in block.get('segments') or [{'text':block['text']}]:
                    if segment.get('figure_ref'):
                        run=field(p,'REF '+_figure_bookmark(segment['figure_ref'])+' \\h',segment['text'])
                    elif segment.get('table_ref'):
                        run=field(p,'REF '+_table_bookmark(segment['table_ref'])+' \\h',segment['text'])
                    else:
                        run=p.add_run(segment['text'])
                    run.bold=bool(segment.get('bold'))
            if any(segment.get('figure_ref') or segment.get('table_ref') for segment in block.get('segments',[])):
                p.paragraph_format.keep_with_next=True
            if style=='bullet':
                p.style=document.styles['List Bullet']
            if style in ('small','meta'):
                for run in p.runs:run.font.size=Pt(10)
            if style=='lead':p.paragraph_format.keep_with_next=True
            if style=='finding':
                p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(8)
            if style=='toc_title':
                for run in p.runs:run.font.size=Pt(17)
                p.paragraph_format.keep_with_next=True
        elif kind=='toc':
            # Native fields recalculate against Word's own pagination; never copy PDF page numbers.
            p=document.add_paragraph()
            field(p,'TOC \\o "1-2" \\h \\z','Open in Word and update the table of contents to display page numbers.')
        elif kind=='reference':
            p=document.add_paragraph()
            link=OxmlElement('w:hyperlink')
            link.set(qn('r:id'),document.part.relate_to(block['url'],RELATIONSHIP_TYPE.HYPERLINK,is_external=True))
            run=OxmlElement('w:r');value=OxmlElement('w:t');value.text=block['text'];run.append(value);link.append(run);p._p.append(link)
        elif kind=='chart':
            p=document.add_paragraph();p.paragraph_format.keep_with_next=True
            run=p.add_run();run.add_picture(BytesIO(base64.b64decode(block['image'])),width=Pt(block.get('width',522)),height=Pt(block['height']))
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after=Pt(0)
            for doc_properties in p._p.xpath('.//wp:docPr'):
                doc_properties.set('descr',block['caption'])
            p=document.add_paragraph(style='Caption')
            if block.get('figure_id') and block.get('figure_number') is not None:
                p.add_run('Figure ')
                identity=bookmark_start(p,_figure_bookmark(block['figure_id']))
                field(p,'SEQ Figure \\* ARABIC',str(block['figure_number']))
                bookmark_end(p,identity)
                p.add_run('. ')
            p.add_run(block['caption'])
            p.paragraph_format.keep_with_next=False
            p.paragraph_format.widow_control=True
            for run in p.runs:run.font.size=Pt(8)
        elif kind=='table':
            if block.get('table_id') and block.get('table_number') is not None:
                p=document.add_paragraph(style='Caption')
                p.add_run('Table ')
                identity=bookmark_start(p,_table_bookmark(block['table_id']))
                field(p,'SEQ Table \\* ARABIC',str(block['table_number']))
                bookmark_end(p,identity)
                p.add_run('. '+block['caption'])
                p.paragraph_format.keep_with_next=True
            emphasis_rows=set(block.get('emphasis_rows',()))
            math_columns=set(block.get('math_columns',()))
            inline_math_columns=set(block.get('inline_math_columns',()))
            table=document.add_table(rows=1,cols=len(block['headers']))
            table.autofit=False
            widths=block.get('widths') or [1/len(block['headers'])]*len(block['headers'])
            for col,width in zip(table.columns,widths):col.width=Pt(522*width)
            header=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(header)
            for row_index,values in enumerate([block['headers']]+block['rows']):
                row=table.rows[0] if row_index==0 else table.add_row()
                if all(len(str(value))<700 for value in values):
                    cant_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant_split)
                for i,(cell,value) in enumerate(zip(row.cells,values)):
                    cell.width=Pt(522*widths[i])
                    if row_index>0 and i in math_columns:
                        cell.text=''
                        _add_equation_runs(cell.paragraphs[0],value)
                    elif row_index>0 and i in inline_math_columns:
                        cell.text=''
                        _add_equation_runs(cell.paragraphs[0],value,mono=False)
                    else:
                        cell.text=str(value)
                    props=cell._tc.get_or_add_tcPr()
                    margins=OxmlElement('w:tcMar')
                    vertical_padding=60 if block.get('compact') else 90
                    for side,twips in (('top',vertical_padding),('bottom',vertical_padding),('left',120),('right',120)):
                        item=OxmlElement('w:'+side);item.set(qn('w:w'),str(twips));item.set(qn('w:type'),'dxa');margins.append(item)
                    props.append(margins)
                    if row_index==0:
                        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EDEDED');props.append(shade)
                    borders=OxmlElement('w:tcBorders');bottom=OxmlElement('w:bottom');bottom.set(qn('w:val'),'single');bottom.set(qn('w:sz'),'3');bottom.set(qn('w:color'),'D0D0D0');borders.append(bottom);props.append(borders)
                    if row_index-1 in emphasis_rows:
                        top=OxmlElement('w:top');top.set(qn('w:val'),'single');top.set(qn('w:sz'),'5');top.set(qn('w:color'),'999999');borders.append(top)
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0
                        p.paragraph_format.keep_with_next=bool(block.get('keep') and row_index<len(block['rows']))
                        if i in block.get('numeric',[]):p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
                        for run in p.runs:run.font.size=Pt(10);run.bold=row_index==0 or row_index-1 in emphasis_rows
            after=document.add_paragraph();after.paragraph_format.space_after=Pt(0);after.paragraph_format.space_before=Pt(0)
            after.paragraph_format.line_spacing=Pt(5)
    output=BytesIO();document.save(output)
    return output.getvalue()


def build_docx(job, *, generated_at=None, include_technical_appendix=True, analysis_name=None):
    from sbepv import technoeconomic_pdf
    from sbepv import technoeconomic_docx_refresh
    report=technoeconomic_pdf.prepare_report(job,generated_at=generated_at,include_technical_appendix=include_technical_appendix,analysis_name=analysis_name)
    return technoeconomic_docx_refresh.refresh_docx_fields(render_docx(report)),technoeconomic_pdf.report_filename(report,'docx')
